from __future__ import annotations

import asyncio
from dataclasses import replace
from datetime import UTC, date, datetime, timedelta
import json
import sqlite3
from types import SimpleNamespace

import pytest
from aiogram.types import InlineKeyboardMarkup, ReplyKeyboardMarkup, ReplyKeyboardRemove

from bot.config import Config
from bot.handlers import voice, work_time as handler
from bot.handlers.decision_callbacks import decision_callback
from bot.handlers.invoice import process_invoice_text
from bot.handlers.state_control import cmd_cancel
from bot.services.active_fsm_guard import handle_active_fsm_text_update
from bot.services.db import init_db
from bot.services.work_time import WorkTimeCandidate, WorkTimeService, day_calculation
from bot.work_time_day_break_migration import audit_day_break_schema, apply_day_break_schema


class State:
    def __init__(self):
        self.current = None
        self.data = {}

    async def get_state(self): return self.current
    async def get_data(self): return dict(self.data)
    async def update_data(self, **kwargs): self.data.update(kwargs)
    async def set_state(self, value): self.current = getattr(value, 'state', value)
    async def clear(self): self.current = None; self.data.clear()


class Telegram:
    id = 9000
    def __init__(self):
        self.messages = []
        self.documents = []
        self.fail_cleanup = False
        self.fail_send = False

    async def edit_message_reply_markup(self, *, chat_id, message_id, reply_markup):
        if self.fail_cleanup: raise RuntimeError('fixture cleanup failure')
        card = next(m for m in self.messages if m.message_id == message_id)
        assert card.chat.id == chat_id
        card.reply_markup = reply_markup

    async def get_file(self, file_id): return SimpleNamespace(file_path='fixture.ogg')
    async def download_file(self, file_path, destination):
        destination.parent.mkdir(parents=True, exist_ok=True)
        destination.write_bytes(b'fake audio')


class Message:
    def __init__(self, bot, text='', actor=1001):
        self.bot, self.text = bot, text
        self.from_user = SimpleNamespace(id=actor, is_bot=False)
        self.chat = SimpleNamespace(id=actor, type='private')
        self.message_id = 77
        self.reply_markup = None

    async def answer(self, text, **kwargs):
        if self.bot.fail_send: raise RuntimeError('fixture send failure')
        card = Message(self.bot, text, self.chat.id)
        card.from_user = SimpleNamespace(id=self.bot.id, is_bot=True)
        card.message_id = len(self.bot.messages) + 100
        card.reply_markup = kwargs.get('reply_markup')
        self.bot.messages.append(card)
        return card

    async def answer_document(self, document, caption=None, **kwargs):
        self.bot.documents.append(document)
        return await self.answer(caption or '', **kwargs)


class Callback:
    def __init__(self, card, *, actor=1001, token=None):
        self.message, self.from_user = card, SimpleNamespace(id=actor)
        self.data = token or card.reply_markup.inline_keyboard[0][0].callback_data
        self.answers = []

    async def answer(self, text=None, **kwargs): self.answers.append(text)


@pytest.fixture
def setup(tmp_path, monkeypatch):
    config = Config(bot_token='test', openai_api_key='sk-test', openai_stt_model='whisper-1',
        openai_llm_model='gpt-4o', debug_invoice_transparency=False, db_path=tmp_path/'time.db',
        storage_dir=tmp_path, allowed_telegram_user_ids=frozenset({1001}), work_time_day_break_enabled=True)
    init_db(config.db_path)
    now = datetime(2026, 10, 20, 17, 0)
    monkeypatch.setattr('bot.services.work_time.work_time_local_now', lambda value=None: value or now)
    monkeypatch.setattr(handler, 'work_time_local_now', lambda value=None: value or now)
    service = WorkTimeService(config.db_path)
    service.save_lunch_break_settings(telegram_id=1001, enabled=True, minutes=30)
    model(monkeypatch)
    return config, service, Telegram()


def seed(service, *, work_date='2026-10-20', duration=False, closed=True):
    service.open_day(telegram_id=1001, now=datetime.fromisoformat(work_date+'T08:00'))
    if closed:
        return service.close_open_day(telegram_id=1001, end_datetime=datetime.fromisoformat(work_date+'T17:00'),
            duration_minutes=300 if duration else None).day
    return service.get_open_day(telegram_id=1001)


def model(monkeypatch, *, slots=None, action='update_work_time_lunch_break', close_mode='close_now'):
    payloads = []
    class Model:
        def __init__(self, **kwargs): self.chat = SimpleNamespace(completions=self)
        async def create(self, **kwargs):
            payload = json.loads(kwargs['messages'][1]['content'])
            payloads.append(payload)
            context = payload.get('context_name', '')
            if context == 'top_level_action':
                value = {'canonical_action': action, 'routing_kind': 'business_action', 'slots': slots or {}}
            elif context == 'work_time_day_break_slots':
                value = slots or {'scope':'unknown'}
            elif context in {'work_time_slot_extraction', 'work_time_open_slot_extraction'}:
                value = {'canonical':'work_time_entry', 'mode':close_mode, 'date':'2026-10-20',
                         'start_time':'07:00' if close_mode == 'open_at_time' else None,
                         'end_time':'17:00' if close_mode == 'close_at_time' else None, 'duration_minutes':None}
            else:
                value = {'canonical':'pass_through', 'canonical_action':'pass_through'}
            return SimpleNamespace(choices=[SimpleNamespace(message=SimpleNamespace(content=json.dumps(value)))])
    monkeypatch.setattr('bot.services.work_time.AsyncOpenAI', Model)
    monkeypatch.setattr('bot.services.semantic_action_resolver.AsyncOpenAI', Model)
    return payloads


def count_events(config):
    with sqlite3.connect(config.db_path) as conn:
        return conn.execute("SELECT count(*) FROM work_time_events WHERE event_type='lunch_break_correction'").fetchone()[0]


@pytest.mark.parametrize('explicit', [False, True])
def test_public_close_finishes_idle_and_optional_button_does_not_start_flow(setup, monkeypatch, explicit):
    config, service, bot = setup
    seed(service, closed=False)
    model(monkeypatch, action='close_work_day', close_mode='close_at_time' if explicit else 'close_now')
    state = State()
    async def journey():
        text = 'Zavri deň o 17:00' if explicit else 'Zavri deň teraz'
        await process_invoice_text(message=Message(bot,text), state=state, config=config, invoice_text=text)
        if explicit:
            assert state.current == handler.WorkTimeStates.waiting_close_preview_confirm.state
            assert service.get_open_day(telegram_id=1001) is not None
            await handler.work_time_close_preview_confirm(Message(bot,'Schváliť'), state, config)
        assert state.current is None and state.data == {}
        assert service.get_open_day(telegram_id=1001) is None
        summary, card = bot.messages[-2:]
        assert 'Trvanie: 9:00' in summary.text and '30 min' in summary.text and '8:30' in summary.text
        assert isinstance(summary.reply_markup, ReplyKeyboardRemove)
        assert card.reply_markup.inline_keyboard[0][0].text == 'Upraviť dnešnú prestávku'
        assert len(card.reply_markup.inline_keyboard[0][0].callback_data.encode()) <= 64
        assert count_events(config) == 0
    asyncio.run(journey())


@pytest.mark.parametrize('entry', ['button','text','voice'])
def test_public_day_zero_break_preserves_default_other_day_and_report(setup, monkeypatch, entry):
    config, service, bot = setup
    seed(service, work_date='2026-10-19')
    original = seed(service)
    payloads = model(monkeypatch, slots={'scope':'day','work_date':'2026-10-20','minutes':0})
    state = State()
    async def journey():
        message = Message(bot, 'Dnes som nemala prestávku, oprav ju na nulu')
        if entry == 'button':
            await handler.send_closed_day_summary(message, config, original)
            card = bot.messages[-1]
            await handler.work_time_day_break_callback(Callback(card), state, config)
            assert card.reply_markup is None
            assert state.current == handler.WorkTimeStates.waiting_day_break_minutes.state
            await handler.work_time_day_break_input(Message(bot,'0'), state, config)
        elif entry == 'voice':
            async def stt(*a, **k): return message.text
            monkeypatch.setattr(voice, 'transcribe_audio', stt)
            message.voice = SimpleNamespace(file_id='fixture')
            await voice.handle_voice(message, bot, config, state)
        else:
            await process_invoice_text(message=message, state=state, config=config, invoice_text=message.text)
        assert state.current == handler.WorkTimeStates.waiting_day_break_confirm.state
        assert service.get_day(telegram_id=1001,work_date=original.work_date) == original
        assert '30 → 0' in bot.messages[-1].text
        await handler.work_time_day_break_confirm(Message(bot,'Schváliť'), state, config)
        saved = service.get_day(telegram_id=1001, work_date=original.work_date)
        assert saved.id == original.id and saved.start_time == '08:00' and saved.end_time == '17:00'
        assert saved.lunch_break_minutes_override == 0 and saved.total_minutes == 540
        assert service.get_lunch_break_settings(telegram_id=1001).minutes == 30
        days = service.list_days_for_month(telegram_id=1001,year=2026,month=10)
        assert sum(service.report_net_minutes(d,settings=service.get_lunch_break_settings(telegram_id=1001)) for d in days) == 1050
        report = service.generate_monthly_report(telegram_id=1001,year=2026,month=10,output_dir=config.storage_dir/'reports')
        from openpyxl import load_workbook
        sheet = load_workbook(report.report_path).active
        assert any(row[0].value == '20.10.2026' and row[3].value == timedelta(minutes=540) for row in sheet.iter_rows())
        assert sheet['D35'].value == timedelta(minutes=1050)
        assert state.current is None and count_events(config) == 1
        await handler.work_time_day_break_confirm(Message(bot,'Schváliť'), state, config)
        assert count_events(config) == 1
    asyncio.run(journey())
    if entry != 'button': assert len([p for p in payloads if p.get('context_name') == 'top_level_action']) == 1


@pytest.mark.parametrize('change_after', ['row','settings','event'])
def test_stale_snapshot_no_overwrite_and_new_approval_needed(setup, change_after):
    config, service, bot = setup
    original = seed(service)
    state = State()
    async def journey():
        await handler.start_update_work_time_lunch_break(Message(bot), state, config, text='',
            break_slots={'scope':'day','work_date':original.work_date,'minutes':0})
        if change_after == 'settings': service.save_lunch_break_settings(telegram_id=1001, enabled=True,minutes=60)
        else:
            with sqlite3.connect(config.db_path) as conn:
                if change_after == 'row': conn.execute('UPDATE work_time_days SET note=? WHERE id=?',('changed',original.id))
                else: conn.execute("INSERT INTO work_time_events(work_time_day_id,telegram_id,event_type) VALUES(?,1001,'fixture')",(original.id,))
        await handler.work_time_day_break_confirm(Message(bot,'Schváliť'), state, config)
        assert count_events(config) == 0
        assert state.current == handler.WorkTimeStates.waiting_day_break_confirm.state
        await handler.work_time_day_break_confirm(Message(bot,'Schváliť'), state, config)
        assert count_events(config) == 1
    asyncio.run(journey())


@pytest.mark.parametrize('exit', ['cancel','command','edit','unknown','expired'])
def test_child_controls_cleanup_and_preserve_saved_closing(setup, exit):
    config, service, bot = setup
    day = seed(service)
    state = State()
    async def journey():
        await handler.start_update_work_time_lunch_break(Message(bot), state, config, text='',
            break_slots={'scope':'day','work_date':day.work_date,'minutes':0})
        if exit == 'expired':
            state.data['work_time_day_break_preview_at']=(datetime.now(UTC)-timedelta(minutes=16)).isoformat()
        if exit == 'command': await cmd_cancel(Message(bot,'/cancel'), state, config)
        else: await handler.work_time_day_break_confirm(Message(bot,{'cancel':'Zrušiť','edit':'Upraviť','unknown':'???'}.get(exit,'Schváliť')),state,config)
        assert service.get_day(telegram_id=1001,work_date=day.work_date) == day
        if exit == 'unknown':
            assert state.current == handler.WorkTimeStates.waiting_day_break_confirm.state
            assert isinstance(bot.messages[-1].reply_markup,ReplyKeyboardMarkup)
        elif exit == 'edit':
            assert state.current == handler.WorkTimeStates.waiting_day_break_minutes.state
            assert isinstance(bot.messages[-1].reply_markup,ReplyKeyboardRemove)
        else:
            assert state.current is None
            assert any(isinstance(m.reply_markup,ReplyKeyboardRemove) for m in bot.messages[1:])
        assert count_events(config)==0
    asyncio.run(journey())


@pytest.mark.parametrize('case', ['active','stale','foreign','malformed','missing','cleanup_failure','disabled'])
def test_launcher_ownership_staleness_and_cleanup(setup, case, caplog):
    config, service, bot = setup
    day=seed(service)
    state=State()
    async def journey():
        await handler.send_closed_day_summary(Message(bot),config,day)
        card=bot.messages[-1]
        token=card.reply_markup.inline_keyboard[0][0].callback_data
        if case=='active': await state.set_state('Other:waiting_value')
        if case=='stale': service.save_lunch_break_settings(telegram_id=1001,enabled=True,minutes=60)
        if case=='missing':
            with sqlite3.connect(config.db_path) as conn: conn.execute('DELETE FROM work_time_days WHERE id=?',(day.id,))
        if case=='cleanup_failure': bot.fail_cleanup=True
        selected_config=replace(config,work_time_day_break_enabled=False) if case=='disabled' else config
        await handler.work_time_day_break_callback(Callback(card,actor=1002 if case=='foreign' else 1001,
            token='wtb:bad' if case=='malformed' else token),state,selected_config)
        assert count_events(config)==0
        if case in {'active','foreign','malformed','missing'}: assert card.reply_markup is not None
        elif case=='cleanup_failure':
            assert state.current==handler.WorkTimeStates.waiting_day_break_minutes.state
            assert card.reply_markup is not None and 'cleanup failed' in caplog.text
        else: assert card.reply_markup is None and state.current is None
        if case=='active': assert state.current=='Other:waiting_value'
    asyncio.run(journey())


def test_day_exception_survives_later_public_arrival_correction_and_default_change(setup, monkeypatch):
    config,service,bot=setup
    day=seed(service)
    change=service.prepare_day_break_change(telegram_id=1001,work_date=day.work_date,minutes=0)
    assert service.apply_day_break_change(telegram_id=1001,change=change).ok
    model(monkeypatch,action='open_work_day',close_mode='open_at_time')
    state=State()
    async def journey():
        text='Зміни сьогодні прихід на роботу на сьому годину'
        await process_invoice_text(message=Message(bot,text),state=state,config=config,invoice_text=text)
        assert state.current==handler.WorkTimeStates.waiting_open_preview_confirm.state
        assert '08:00' in bot.messages[-1].text and '07:00' in bot.messages[-1].text
        await handler.work_time_open_preview_confirm(Message(bot,'Schváliť'),state,config)
        saved=service.get_day(telegram_id=1001,work_date=day.work_date)
        assert saved.start_time=='07:00' and saved.end_time=='17:00' and saved.status=='closed'
        assert saved.lunch_break_minutes_override==0 and saved.total_minutes==600
        service.save_lunch_break_settings(telegram_id=1001,enabled=True,minutes=60)
        assert service.report_net_minutes(saved,settings=service.get_lunch_break_settings(telegram_id=1001))==600
        assert state.current is None
    asyncio.run(journey())


@pytest.mark.parametrize('value', [-1,181,True,540])
def test_invalid_break_has_no_service_write(setup,value):
    config,service,_=setup
    day=seed(service)
    assert service.prepare_day_break_change(telegram_id=1001,work_date=day.work_date,minutes=value) is None
    assert count_events(config)==0


def test_fixed_net_exclusion_and_old_inline_cannot_approve_new_preview(setup):
    config,service,bot=setup
    day=seed(service,duration=True)
    assert service.prepare_day_break_change(telegram_id=1001,work_date=day.work_date,minutes=0) is None
    state=State()
    async def journey():
        await handler.send_closed_day_summary(Message(bot),config,day)
        assert not any(isinstance(m.reply_markup,InlineKeyboardMarkup) for m in bot.messages)
        assert 'Potvrdené čisté' in bot.messages[-1].text
        await state.set_state(handler.WorkTimeStates.waiting_day_break_confirm)
        cb=Callback(bot.messages[-1],token='decision:approve')
        await decision_callback(cb,state,config)
        assert count_events(config)==0
    asyncio.run(journey())


def test_explicit_schema_migration_preserves_legacy_values_and_no_startup_alter(setup,tmp_path):
    config,service,_=setup
    day=seed(service)
    with sqlite3.connect(config.db_path) as conn: conn.execute('ALTER TABLE work_time_days DROP COLUMN lunch_break_minutes_override')
    before=audit_day_break_schema(config.db_path)
    init_db(config.db_path)
    assert not service.day_break_schema_ready()
    assert service.get_day(telegram_id=1001,work_date=day.work_date)==day
    assert audit_day_break_schema(config.db_path)['fingerprint']==before['fingerprint']
    with pytest.raises(RuntimeError,match='baseline_changed'):
        apply_day_break_schema(config.db_path,expected_fingerprint='wrong',backup_dir=tmp_path/'backups',service_stopped=True)
    assert not (tmp_path/'backups').exists()
    with pytest.raises(RuntimeError,match='writer_freeze'):
        apply_day_break_schema(config.db_path,expected_fingerprint=before['fingerprint'],backup_dir=tmp_path/'backups')
    result=apply_day_break_schema(config.db_path,expected_fingerprint=before['fingerprint'],backup_dir=tmp_path/'backups',service_stopped=True)
    assert result['business_digest']==before['business_digest'] and result['integrity']=='ok'
    assert audit_day_break_schema(__import__('pathlib').Path(result['backup_path']))['fingerprint']==before['fingerprint']
    assert service.get_day(telegram_id=1001,work_date=day.work_date)==day
    with sqlite3.connect(config.db_path) as conn:
        assert conn.execute('SELECT lunch_break_minutes_override FROM work_time_days').fetchall()==[(None,)]


def test_scope_clarification_then_voice_minutes_and_shared_voice_confirmation(setup,monkeypatch):
    config,service,bot=setup
    day=seed(service)
    state=State()
    model(monkeypatch,slots={'scope':'unknown','minutes':None,'work_date':None})
    async def journey():
        await process_invoice_text(message=Message(bot,'Zmeň prestávku'),state=state,config=config,invoice_text='Zmeň prestávku')
        assert state.current==handler.WorkTimeStates.waiting_lunch_scope.state
        model(monkeypatch,slots={'scope':'day','work_date':day.work_date,'minutes':None})
        await handler.work_time_day_break_input(Message(bot,'iba dnes'),state,config)
        assert state.current==handler.WorkTimeStates.waiting_day_break_minutes.state
        model(monkeypatch,slots={'scope':'day','work_date':day.work_date,'minutes':15})
        async def transcript(*a,**k): return 'pätnásť minút'
        monkeypatch.setattr(voice,'transcribe_audio',transcript)
        message=Message(bot); message.voice=SimpleNamespace(file_id='fixture')
        await voice.handle_voice(message,bot,config,state)
        assert state.current==handler.WorkTimeStates.waiting_day_break_confirm.state
        async def approve(*a,**k): return 'Schváliť'
        monkeypatch.setattr(voice,'transcribe_audio',approve)
        await voice.handle_voice(message,bot,config,state)
        saved=service.get_day(telegram_id=1001,work_date=day.work_date)
        assert saved.total_minutes==525 and saved.lunch_break_minutes_override==15
        assert state.current is None and count_events(config)==1
        assert not list(config.storage_dir.rglob('*.ogg'))
    asyncio.run(journey())


@pytest.mark.parametrize('case',['menu','start','timeout','help'])
def test_navigation_cleanup_preserves_day_and_leaves_child_flow(setup,case):
    config,service,bot=setup
    day=seed(service)
    state=State()
    async def journey():
        await handler.start_update_work_time_lunch_break(Message(bot),state,config,text='',
            break_slots={'scope':'day','work_date':day.work_date,'minutes':0})
        if case=='help': await handler.cmd_dochadzka(Message(bot,'/dochadzka'),state,config)
        else:
            from bot.services.active_fsm_guard import ACTIVE_FSM_LAST_ACTIVITY_AT_KEY, expire_active_fsm_state
            if case=='timeout': await expire_active_fsm_state(message=Message(bot),state=state,config=config)
            else: await handle_active_fsm_text_update(message=Message(bot,'/'+case),state=state,config=config,text='/'+case,input_channel='text')
        assert state.current!=handler.WorkTimeStates.waiting_day_break_confirm.state
        assert not state.data.get(handler.DAY_BREAK_CHANGE_KEY)
        assert service.get_day(telegram_id=1001,work_date=day.work_date)==day
        assert any(isinstance(m.reply_markup,ReplyKeyboardRemove) for m in bot.messages[1:])
        assert count_events(config)==0
    asyncio.run(journey())


@pytest.mark.parametrize('case',['workspace','membership','actor'])
def test_authority_change_between_preview_and_confirmation_cannot_write(setup,case):
    config,_,bot=setup
    with sqlite3.connect(config.db_path) as conn:
        conn.execute("INSERT INTO authorized_users(telegram_id,created_at) VALUES(1001,CURRENT_TIMESTAMP)")
        for key in ['one','two']:
            conn.execute("INSERT INTO workspace VALUES(?,?,?,?,?,CURRENT_TIMESTAMP,CURRENT_TIMESTAMP)",(key,key,key,key,'active'))
            conn.execute("INSERT INTO workspace_membership VALUES(?,1001,'owner','active',CURRENT_TIMESTAMP,CURRENT_TIMESTAMP)",(key,))
        conn.execute("INSERT INTO active_workspace_selection VALUES(1001,'one',CURRENT_TIMESTAMP)")
    service=WorkTimeService(config.db_path,workspace_id='one')
    other=WorkTimeService(config.db_path,workspace_id='two')
    day=seed(service); other_day=seed(other)
    state=State()
    async def journey():
        await handler.start_update_work_time_lunch_break(Message(bot),state,config,text='',
            break_slots={'scope':'day','work_date':day.work_date,'minutes':0})
        with sqlite3.connect(config.db_path) as conn:
            if case=='workspace': conn.execute("UPDATE active_workspace_selection SET workspace_id='two' WHERE telegram_id=1001")
            if case=='membership': conn.execute("UPDATE workspace_membership SET status='inactive' WHERE workspace_id='one'")
        await handler.work_time_day_break_confirm(Message(bot,'Schváliť',actor=1002 if case=='actor' else 1001),state,config)
        assert service.get_day(telegram_id=1001,work_date=day.work_date)==day
        assert other.get_day(telegram_id=1001,work_date=day.work_date)==other_day
        assert count_events(config)==0
    asyncio.run(journey())


def test_disabled_entry_and_later_boundary_keep_saved_exception(setup):
    config,service,bot=setup
    day=seed(service)
    change=service.prepare_day_break_change(telegram_id=1001,work_date=day.work_date,minutes=15)
    saved=service.apply_day_break_change(telegram_id=1001,change=change).day
    async def journey():
        await handler.send_closed_day_summary(Message(bot),replace(config,work_time_day_break_enabled=False),saved)
        assert len(bot.messages)==1 and isinstance(bot.messages[0].reply_markup,ReplyKeyboardRemove)
        assert '15 min' in bot.messages[0].text
    asyncio.run(journey())
    assert service.prepare_boundary_change(telegram_id=1001,work_date=day.work_date,field='end_time',value='08:10') is None
    change=service.prepare_boundary_change(telegram_id=1001,work_date=day.work_date,field='start_time',value='07:00')
    result=service.apply_boundary_change(telegram_id=1001,change=change)
    assert result.day.total_minutes==585 and result.day.lunch_break_minutes_override==15
    # Older readers see the compatibility projection; rollout rollback must still keep compatible writers.
    with sqlite3.connect(config.db_path) as conn:
        assert conn.execute('SELECT net_work_minutes_override FROM work_time_days WHERE id=?',(saved.id,)).fetchone()[0]==585


def test_callback_does_not_take_over_flow_started_during_telegram_cleanup(setup):
    config,service,bot=setup
    day=seed(service); state=State()
    async def journey():
        await handler.send_closed_day_summary(Message(bot),config,day)
        card=bot.messages[-1]
        original=bot.edit_message_reply_markup
        async def cleanup(**kwargs):
            await original(**kwargs)
            await state.set_state('Other:waiting_value')
        bot.edit_message_reply_markup=cleanup
        await handler.work_time_day_break_callback(Callback(card),state,config)
        assert state.current=='Other:waiting_value' and count_events(config)==0
    asyncio.run(journey())


def test_unknown_or_cancel_during_slot_extraction_cannot_resurrect_flow(setup,monkeypatch):
    config,service,bot=setup
    day=seed(service); state=State()
    async def journey():
        await handler.start_update_work_time_lunch_break(Message(bot),state,config,text='',
            break_slots={'scope':'day','work_date':day.work_date})
        async def delayed(**kwargs):
            await state.clear()
            return {'scope':'day','work_date':day.work_date,'minutes':0}
        monkeypatch.setattr(handler,'resolve_day_break_slots',delayed)
        await handler.work_time_day_break_input(Message(bot,'bez prestávky'),state,config)
        assert state.current is None and count_events(config)==0
    asyncio.run(journey())


def test_truth_help_explain_feature_gate_and_no_daily_prompt():
    from bot.services.product_truth import get_safe_answer_payload
    from bot.services.info_help import build_product_truth_guidance
    payload=get_safe_answer_payload('work_time_tracking')
    assert payload['product_status']=='partial'
    text=' '.join(payload['current_limitations'])
    assert 'WORK_TIME_DAY_BREAK_ENABLED' in text and 'stays idle' in text and 'Fixed net-duration' in text
    answer=build_product_truth_guidance(user_input_text='Ako zmením prestávku?')
    assert 'Upraviť dnešnú prestávku' in answer
    assert 'obvykle nastavenie a ostatne dni zostanu' in answer


def test_ignore_button_then_public_report_and_help_do_not_enter_break_edit(setup,monkeypatch):
    config,service,bot=setup
    day=seed(service); state=State()
    model(monkeypatch,action='generate_work_time_report',slots={'period':{'type':'month','month':10,'year':2026}})
    async def journey():
        await handler.send_closed_day_summary(Message(bot),config,day)
        card=bot.messages[-1]
        await process_invoice_text(message=Message(bot,'Ukáž výkaz za október'),state=state,config=config,invoice_text='Ukáž výkaz za október')
        assert state.current is None and len(bot.documents)==1
        assert isinstance(card.reply_markup,InlineKeyboardMarkup)
        assert count_events(config)==0 and service.get_day(telegram_id=1001,work_date=day.work_date)==day
        await process_invoice_text(message=Message(bot,'Ako zmením prestávku?'),state=state,
            config=replace(config,openai_api_key=None),invoice_text='Ako zmením prestávku?')
        assert state.current is None and count_events(config)==0
        assert service.get_day(telegram_id=1001,work_date=day.work_date)==day
    asyncio.run(journey())


def test_committed_correction_survives_result_delivery_failure(setup,caplog):
    config,service,bot=setup
    day=seed(service); state=State()
    async def journey():
        await handler.start_update_work_time_lunch_break(Message(bot),state,config,text='',
            break_slots={'scope':'day','work_date':day.work_date,'minutes':0})
        bot.fail_send=True
        await handler.work_time_day_break_confirm(Message(bot,'Schváliť'),state,config)
        assert state.current is None
        assert service.get_day(telegram_id=1001,work_date=day.work_date).lunch_break_minutes_override==0
        assert count_events(config)==1 and 'delivery failed' in caplog.text
    asyncio.run(journey())


@pytest.mark.parametrize('slots', [
    {'scope':{},'minutes':0}, {'scope':'day','work_date':'bad-date','minutes':0},
    {'scope':'day','work_date':'2026-10-21','minutes':0},
    {'scope':'day','work_date':'2026-10-20','minutes':True},
])
def test_invalid_extracted_slots_clarify_without_write(setup,slots):
    config,service,bot=setup
    day=seed(service);state=State()
    asyncio.run(handler.start_update_work_time_lunch_break(Message(bot),state,config,text='',break_slots=slots))
    assert state.current in {handler.WorkTimeStates.waiting_lunch_scope.state,
        handler.WorkTimeStates.waiting_day_break_date.state,handler.WorkTimeStates.waiting_day_break_minutes.state}
    assert service.get_day(telegram_id=1001,work_date=day.work_date)==day and count_events(config)==0


def test_day_break_rederives_gross_from_clocks_without_silent_legacy_backfill(setup):
    config,service,_=setup
    day=seed(service)
    with sqlite3.connect(config.db_path) as conn:
        conn.execute('UPDATE work_time_days SET gross_minutes=777,total_minutes=747 WHERE id=?',(day.id,))
    original=service.get_day(telegram_id=1001,work_date=day.work_date)
    assert service.report_net_minutes(original,settings=service.get_lunch_break_settings(telegram_id=1001))==747
    change=service.prepare_day_break_change(telegram_id=1001,work_date=day.work_date,minutes=0)
    assert change.proposed.gross_minutes==540 and change.proposed.total_minutes==540
    # Preparation is read-only, including an inconsistent legacy gross field.
    assert service.get_day(telegram_id=1001,work_date=day.work_date)==original
    saved=service.apply_day_break_change(telegram_id=1001,change=change).day
    assert saved.net_work_minutes_override==540 and day_calculation(saved,None)==(540,0,540)


def test_existing_workspace_copy_preserves_explicit_zero_and_compatibility_net(setup,tmp_path):
    from bot.services.multi_workspace_migration_apply import MultiWorkspaceMigrationManager
    config,service,_=setup
    day=seed(service)
    change=service.prepare_day_break_change(telegram_id=1001,work_date=day.work_date,minutes=0)
    assert service.apply_day_break_change(telegram_id=1001,change=change).ok
    target_path=tmp_path/'copied.db'
    init_db(target_path)
    manager=MultiWorkspaceMigrationManager(db_path=config.db_path,storage_root=config.storage_dir)
    candidate=SimpleNamespace(telegram_id=1001,supplier_id=1,workspace_id='copied')
    with sqlite3.connect(config.db_path) as source, sqlite3.connect(target_path) as target:
        source.row_factory=sqlite3.Row
        manager._copy_table(source,target,'work_time_days',SimpleNamespace(candidates=[candidate]))
        target.commit()
    copied=WorkTimeService(target_path,workspace_id='copied').get_day(telegram_id=1001,work_date=day.work_date)
    assert copied.lunch_break_minutes_override==0 and copied.net_work_minutes_override==540
    assert copied.start_time==day.start_time and copied.end_time==day.end_time
